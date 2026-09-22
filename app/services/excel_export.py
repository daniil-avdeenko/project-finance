"""
Экспорт данных в Excel (.xlsx) через openpyxl.
"""

import logging
from io import BytesIO

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.currency_service import convert_with_rates

logger = logging.getLogger(__name__)

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Стиль заголовков
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def _setup_sheet(ws, headers: list[str]) -> None:
    """Записывает заголовки, применяет стили и включает автофильтр."""
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    ws.freeze_panes = "A2"  # закрепляем строку заголовков
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _autosize_columns(
    ws,
    *,
    min_width: int = 10,
    max_width: int = 80,
) -> None:
    """
    Подбирает ширину колонок по максимальной длине значения.
    """
    for col_idx, column in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        width = max(min_width, min(max_length + 6, max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def make_xlsx_response(workbook: Workbook, filename: str) -> Response:
    """Сериализует Workbook в HTTP-ответ с правильными заголовками."""
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype=MIME_XLSX,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _new_workbook() -> Workbook:
    """Убираем дефолтный лист — будем создавать свой с именем."""
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def build_projects_workbook(projects: list) -> Workbook:
    """Workbook с листом «Проекты»."""
    wb = _new_workbook()
    ws = wb.create_sheet("Проекты")

    _setup_sheet(
        ws,
        [
            "ID",
            "Название",
            "Описание",
            "Доходы (₽)",
            "Расходы (₽)",
            "Прибыль (₽)",
            "Рентабельность (%)",
            "Кол-во сотрудников",
            "Кол-во транзакций",
        ],
    )

    for p in projects:
        ws.append(
            [
                p.id,
                p.name,
                p.description or "",
                round(p.total_income, 2),
                round(p.total_expense, 2),
                round(p.profit, 2),
                p.profitability,
                len(p.employee_roles),
                p.transactions.filter_by(is_deleted=False).count(),
            ]
        )

    _autosize_columns(ws)
    return wb


def build_transactions_workbook(transactions: list, rates_map: dict) -> Workbook:
    """
    Workbook с листом «Транзакции».

    rates_map — предзагруженные курсы ЦБ (см. get_rates_map).
    Сумма в рублях считается в памяти, без N+1 к currency_rates.
    """
    wb = _new_workbook()
    ws = wb.create_sheet("Транзакции")

    _setup_sheet(
        ws,
        [
            "ID",
            "Дата",
            "Проект",
            "Тип",
            "Категория",
            "Сумма",
            "Валюта",
            "Сумма (₽)",
            "Описание",
        ],
    )

    # Справочники категорий — по одному запросу на тип, а не на строку
    from app.models import ExpenseCategory, IncomeCategory

    income_cats = {c.id: c.name for c in IncomeCategory.query.all()}
    expense_cats = {c.id: c.name for c in ExpenseCategory.query.all()}

    for t in transactions:
        on_date = t.date.date() if t.date else None
        rub = convert_with_rates(t.amount, t.currency or "RUB", on_date, rates_map)
        category = (
            income_cats.get(t.category_id, "")
            if t.type == "income"
            else expense_cats.get(t.category_id, "")
        )

        ws.append(
            [
                t.id,
                t.date.strftime("%d.%m.%Y %H:%M") if t.date else "",
                t.project.name if t.project else "",
                "Доход" if t.type == "income" else "Расход",
                category,
                round(t.amount, 2),
                t.currency or "RUB",
                rub,
                t.description or "",
            ]
        )

    _autosize_columns(ws)
    return wb


def build_employees_workbook(employees: list) -> Workbook:
    """Workbook с листом «Сотрудники»."""
    wb = _new_workbook()
    ws = wb.create_sheet("Сотрудники")

    _setup_sheet(ws, ["ID", "ФИО", "Телефон", "Email", "Проекты и роли"])

    for emp in employees:
        roles = "; ".join(f"{er.project.name} — {er.role}" for er in emp.project_roles)
        ws.append([emp.id, emp.name, emp.phone or "", emp.email or "", roles])

    _autosize_columns(ws)
    return wb
