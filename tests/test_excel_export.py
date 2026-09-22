"""Тесты Excel-экспорта через openpyxl."""

from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import db as _db
from app.models import Project
from app.services.excel_export import (
    build_employees_workbook,
    build_projects_workbook,
    build_transactions_workbook,
    make_xlsx_response,
)

# ============================================================
#   ФИКСТУРЫ
# ============================================================


@pytest.fixture
def fake_project():
    class FakeProject:
        id = 1
        name = "CRM для банка"
        description = "Описание"
        total_income = 100_000.0
        total_expense = 40_000.0
        profit = 60_000.0
        profitability = 60.0

        class _Emp:
            pass

        employee_roles = [_Emp(), _Emp()]

        class _TxQuery:
            def filter_by(self, **kw):
                return self

            def count(self):
                return 25

        transactions = _TxQuery()

    return FakeProject()


@pytest.fixture
def fake_transaction():
    class FakeProject:
        name = "CRM для банка"

    class FakeTransaction:
        id = 42
        date = datetime(2026, 9, 15, 12, 30, tzinfo=UTC)
        project = FakeProject()
        project_id = 1
        type = "income"
        category_id = 1
        amount = 100.0
        currency = "USD"
        description = "Оплата"

    return FakeTransaction()


@pytest.fixture
def fake_employee():
    class FakeProject:
        name = "CRM для банка"

    class FakeEmployeeProject:
        project = FakeProject()
        role = "Разработчик"

    class FakeEmployee:
        id = 1
        name = "Иванов Иван"
        phone = "+79990000000"
        email = "ivanov@test.ru"
        project_roles = [FakeEmployeeProject()]

    return FakeEmployee()


# ============================================================
#   build_projects_workbook
# ============================================================


def test_projects_workbook_has_headers_and_row(app, fake_project):
    with app.app_context():
        wb = build_projects_workbook([fake_project])

    ws = wb["Проекты"]
    headers = [c.value for c in ws[1]]
    assert "ID" in headers
    assert "Название" in headers
    assert "Прибыль (₽)" in headers

    row = [c.value for c in ws[2]]
    assert row[0] == 1
    assert row[1] == "CRM для банка"
    assert row[5] == 60_000.0  # Прибыль
    assert isinstance(row[5], float)  # именно число, не строка


def test_projects_workbook_empty_list(app):
    with app.app_context():
        wb = build_projects_workbook([])

    ws = wb["Проекты"]
    assert ws.max_row == 1  # только заголовки


# ============================================================
#   build_transactions_workbook
# ============================================================


def test_transactions_workbook_uses_rates_map(app, fake_transaction):
    """Сумма в рублях считается через rates_map, без SQL."""

    class FakeRate:
        code = "USD"
        rate_date = date(2026, 9, 15)
        nominal = 1
        rate_rub = 84.50

    # Подсовываем фейковый rate в rates_map напрямую
    rates_map = {"USD": [FakeRate()]}

    with app.app_context():
        wb = build_transactions_workbook([fake_transaction], rates_map)

    ws = wb["Транзакции"]
    row = [c.value for c in ws[2]]

    assert row[0] == 42
    assert row[1] == "15.09.2026 12:30"
    assert row[3] == "Доход"
    assert row[5] == 100.0  # исходная сумма
    assert row[6] == "USD"
    assert row[7] == 8450.0  # конвертированная


def test_transactions_workbook_empty(app):
    with app.app_context():
        wb = build_transactions_workbook([], {})
    ws = wb["Транзакции"]
    assert ws.max_row == 1


# ============================================================
#   build_employees_workbook
# ============================================================


def test_employees_workbook_formats_roles(app, fake_employee):
    with app.app_context():
        wb = build_employees_workbook([fake_employee])

    ws = wb["Сотрудники"]
    row = [c.value for c in ws[2]]
    assert row[0] == 1
    assert row[1] == "Иванов Иван"
    assert "CRM для банка — Разработчик" in row[4]


# ============================================================
#   make_xlsx_response
# ============================================================


def test_make_xlsx_response_headers(app):
    with app.app_context():
        wb = build_projects_workbook([])
        response = make_xlsx_response(wb, "test.xlsx")

    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["Content-Disposition"]
    assert "test.xlsx" in response.headers["Content-Disposition"]
    assert "no-store" in response.headers["Cache-Control"]


def test_make_xlsx_response_body_is_valid_xlsx(app):
    """Скачанный файл — валидный .xlsx."""
    with app.app_context():
        wb = build_projects_workbook([])
        response = make_xlsx_response(wb, "test.xlsx")

    reopened = load_workbook(BytesIO(response.get_data()), read_only=True)
    assert "Проекты" in reopened.sheetnames


# ============================================================
#   Роуты
# ============================================================


def test_export_projects_xlsx_route(auth_client, app):
    with app.app_context():
        _db_proj = Project(name="P1")

        _db.session.add(_db_proj)
        _db.session.commit()

    response = auth_client.get("/projects/export/xlsx")
    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_export_transactions_xlsx_route(auth_client):
    response = auth_client.get("/transactions/export/xlsx")
    assert response.status_code == 200


def test_export_employees_xlsx_route(auth_client):
    response = auth_client.get("/employees/export/xlsx")
    assert response.status_code == 200


def test_export_responses_have_distinct_filenames(auth_client):
    """Каждый экспорт отдаёт файл со своим именем и запрещает кеш."""
    endpoints = [
        ("/projects/export/xlsx", "projects_export.xlsx"),
        ("/transactions/export/xlsx", "transactions_export.xlsx"),
        ("/employees/export/xlsx", "employees_export.xlsx"),
    ]
    for url, expected_filename in endpoints:
        response = auth_client.get(url)
        cd = response.headers["Content-Disposition"]
        assert expected_filename in cd, f"{url}: ожидали {expected_filename}, получили {cd}"
        assert "no-store" in response.headers.get("Cache-Control", ""), url
